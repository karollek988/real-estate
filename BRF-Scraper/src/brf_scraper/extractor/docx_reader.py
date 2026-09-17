"""DOCX text extraction for annual reports, producing the same PDFDocument
shape pdf_reader.py produces. Every downstream extractor
(financial_extractor, property_extractor, validation) already reads
PDFDocument.full_text / .pages[].tables / .search_text() / .find_in_tables()
— a Word document needs no extractor changes, only a way to arrive at that
same object. Docx has no reliable page concept once rendered, so the whole
document becomes a single "page".
"""
from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from brf_scraper.utils.logging import get_logger

from .pdf_reader import PageText, PDFDocument

logger = get_logger(__name__)


def read_docx(path: str | Path) -> PDFDocument:
    """Extract paragraph and table text from a .docx file."""
    path = str(path)
    doc = PDFDocument(path=path, total_pages=1)

    try:
        docx = DocxDocument(path)
        paragraphs = [p.text for p in docx.paragraphs if p.text.strip()]

        tables: list[list[list[str]]] = []
        for table in docx.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            tables.append(rows)
            # Regex-based extractors scan full_text, not .tables, so fold
            # each table's cells into the plain-text body too — a balance
            # sheet or loan table's label/value pairs need to be findable
            # either way.
            paragraphs.extend(" ".join(cell for cell in row if cell) for row in rows)

        doc.pages.append(PageText(page_number=1, text="\n".join(paragraphs), tables=tables))
    except Exception as e:
        logger.error("docx_read_failed", path=path, error=str(e))

    if doc.is_text_based:
        logger.info("docx_text_extracted", path=path, total_chars=sum(p.char_count for p in doc.pages))
    else:
        logger.warning("docx_no_text_extracted", path=path)

    return doc
