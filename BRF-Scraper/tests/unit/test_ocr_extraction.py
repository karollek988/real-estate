"""Tests for OCR-backed extraction (extractor/ocr.py, docx_reader.py,
image_reader.py, and pdf_reader.py's scanned-page OCR fallback).

Renders real images/documents and runs them through actual Tesseract
(skipped if tesseract isn't on PATH in this environment, e.g. CI without
the system package installed) rather than mocking OCR, since the thing
worth verifying is that these modules actually produce usable text/fields
from real input, not just that they call a library correctly.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument
from PIL import Image, ImageDraw, ImageFont

from brf_scraper.extractor.docx_reader import read_docx
from brf_scraper.extractor.engine import extract_annual_report
from brf_scraper.extractor.image_reader import read_image
from brf_scraper.extractor.ocr import ocr_image
from brf_scraper.extractor.pdf_reader import read_pdf

# A real (non-bitmap) TrueType font is needed to render legible test images —
# PIL's own load_default() bitmap font is too small/low-res for Tesseract to
# read reliably. None of these paths are guaranteed to exist on any given
# machine (Windows dev boxes, the Linux production container, and CI images
# all ship different fonts by default), so this is a best-effort search
# across the common ones rather than a single hardcoded path.
_FONT_CANDIDATES = [
    "C:/Windows/Fonts/arial.ttf",  # Windows
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",  # Debian/Ubuntu (fonts-dejavu-core)
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",  # Debian/Ubuntu (fonts-liberation)
    "/usr/share/fonts/TTF/DejaVuSans.ttf",  # Arch
    "/System/Library/Fonts/Supplemental/Arial.ttf",  # macOS
]
FONT_PATH = next((p for p in _FONT_CANDIDATES if Path(p).is_file()), None)

pytestmark = pytest.mark.skipif(
    shutil.which("tesseract") is None or FONT_PATH is None,
    reason="tesseract binary or a usable TrueType font is not available in this environment",
)


def _render_text_image(lines: list[str], width: int = 900, height: int = 400) -> Image.Image:
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(FONT_PATH, 32)
    y = 20
    for line in lines:
        draw.text((20, y), line, fill="black", font=font)
        y += 50
    return image


class TestOcrImage:
    def test_extracts_recognizable_swedish_text(self):
        image = _render_text_image(["Utgangspris 4500000 kr", "Boarea 65 kvadratmeter"])
        text = ocr_image(image)
        normalized = text.lower()
        assert "4500000" in normalized.replace(" ", "")
        assert "boarea" in normalized or "65" in normalized


class TestReadImage:
    def test_reads_text_from_image_file(self, tmp_path: Path):
        image = _render_text_image(["Arsredovisning 2025", "Omsattning 1200000 kr"])
        path = tmp_path / "report_page.png"
        image.save(path)

        doc = read_image(path)

        assert doc.is_text_based
        assert "1200000" in doc.full_text.replace(" ", "")


class TestReadDocx:
    def test_extracts_paragraphs_and_tables(self, tmp_path: Path):
        docx = DocxDocument()
        docx.add_paragraph("Arsredovisning for Brf Testgatan 2025")
        table = docx.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Rorelsens intakter"
        table.rows[0].cells[1].text = "1500000"
        table.rows[1].cells[0].text = "Rorelsens kostnader"
        table.rows[1].cells[1].text = "900000"
        path = tmp_path / "annual_report.docx"
        docx.save(path)

        doc = read_docx(path)

        assert doc.is_text_based
        assert "Testgatan" in doc.full_text
        # Table content must also be reachable via find_in_tables() (used by
        # the same financial extractors PDFs go through) and via full_text
        # (used by the regex-based extractors).
        results = doc.find_in_tables("Rorelsens intakter", value_column=1)
        assert results == [(1, "1500000")]
        assert "1500000" in doc.full_text


class TestExtractAnnualReportDocx:
    def test_docx_file_kind_reaches_financial_extraction(self, tmp_path: Path):
        docx = DocxDocument()
        docx.add_paragraph("Arsredovisning for Brf Testgatan, rakenskapsar 2025")
        table = docx.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Nettoomsattning"
        table.rows[0].cells[1].text = "1 500 000 kr"
        path = tmp_path / "annual_report.docx"
        docx.save(path)

        result = extract_annual_report(path, file_kind="docx")

        assert result.is_text_based
        assert result.total_pages == 1


class TestReadPdfOcrFallback:
    def test_scanned_page_falls_back_to_ocr(self, tmp_path: Path):
        # Build a PDF with an image-only page (no embedded text layer at
        # all) -- exactly the "scanned document" case this fallback exists
        # for. pdfplumber's extract_text() must find nothing here; the OCR
        # path is what's actually being verified.
        image = _render_text_image(["Skannad sida i arsredovisningen", "Langfristiga skulder 2500000 kr"])
        image_path = tmp_path / "page.png"
        image.save(image_path)

        pdf_path = tmp_path / "scanned.pdf"
        pdf = fitz.open()
        page = pdf.new_page(width=900, height=400)
        page.insert_image(fitz.Rect(0, 0, 900, 400), filename=str(image_path))
        pdf.save(pdf_path)
        pdf.close()

        doc = read_pdf(pdf_path)

        assert doc.is_text_based, "OCR fallback should have populated page text"
        assert "2500000" in doc.full_text.replace(" ", "")
